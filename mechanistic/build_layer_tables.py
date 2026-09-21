"""Complete 31-layer promoted/suppressed token tables from the ARCHIVED Jacobian tensors.

CPU only. No model, no GPU, no Jacobian recomputation -- the three delta tensors are read
from disk exactly as archived.

Ranking is mechanical and identical for every layer and contrast:
    promoted   = 10 largest mean signed delta
    suppressed = 10 smallest (most negative) mean signed delta
Sign consistency = fraction of the 72 tasks whose delta has the same sign as the mean.

SIGN CONVENTION, carried in every row: positive for A_minus_B => token MORE PROMOTED under
condition A.
SUBWORD CAVEAT, carried in every row: BPE fragments, not words or concepts.
"""
import argparse
import csv
import json
import pathlib

import numpy as np

HERE = pathlib.Path(__file__).resolve().parent
N_TOP = 10
LOCKED = 24
SUBWORD = ("BPE tokenizer fragment, not a word or concept; may be a word piece or a "
           "leading-space variant")
ORDER = ["self_low_2_minus_no_history", "self_high_10_minus_self_low_2",
         "self_low_2_minus_unrelated_low_2"]

ap = argparse.ArgumentParser()
ap.add_argument("--tensor-dir", default=str(HERE / "out/jlens_tokens_raw"))
ap.add_argument("--vocab", default=str(HERE / "out/jlens_tokens_analysis/vocab_strings_repaired.json"))
ap.add_argument("--out-dir", default=str(HERE / "out/jlens_layer_tables"))
A = ap.parse_args()

TD = pathlib.Path(A.tensor_dir)
OUT = pathlib.Path(A.out_dir); OUT.mkdir(parents=True, exist_ok=True)
V = json.loads(pathlib.Path(A.vocab).read_text(encoding="utf-8"))
meta = json.loads((TD / "jlens_tokens_meta.json").read_text(encoding="utf-8"))
LAYERS = meta["our_layers_covered"]


def tokstr(i):
    return V.get(str(i), f"<UNMAPPED_{i}>")


rows = []
per_contrast = {}
for comp in ORDER:
    a, b = comp.split("_minus_")
    arr = np.load(TD / f"delta_{comp}.npy", mmap_mode="r")     # [72, 31, vocab]
    n_tasks = arr.shape[0]
    table = []
    for li, L in enumerate(LAYERS):
        block = np.asarray(arr[:, li, :], dtype=np.float32)     # [72, vocab]
        mean = block.mean(0)
        order = np.argsort(mean)
        sup, pro = order[:N_TOP], order[::-1][:N_TOP]
        for role, ids in (("promoted", pro), ("suppressed", sup)):
            med = np.median(block[:, ids], axis=0)
            sgn = np.sign(mean[ids])
            cons = (np.sign(block[:, ids]) == sgn).mean(0)
            for rank, (tid, m, md, c) in enumerate(zip(ids, mean[ids], med, cons), 1):
                r = {
                    "contrast": comp, "condition_A": a, "condition_B": b,
                    "layer_ours": L, "layer_locked": (L == LOCKED),
                    "role": role, "rank": rank,
                    "token_id": int(tid), "token_string": tokstr(int(tid)),
                    "token_repr": repr(tokstr(int(tid))),
                    "mean_delta_logprob": round(float(m), 6),
                    "median_delta_logprob": round(float(md), 6),
                    "sign_consistency": round(float(c), 4),
                    "n_tasks": n_tasks,
                    "positive_means_more_promoted_under": a,
                    "subword_caveat": SUBWORD,
                }
                rows.append(r); table.append(r)
    per_contrast[comp] = table
    del arr

long_csv = OUT / "jlens_layer_tables_long.csv"
with long_csv.open("w", newline="", encoding="utf-8") as f:
    w = csv.DictWriter(f, fieldnames=list(rows[0]))
    w.writeheader(); w.writerows(rows)
print(f"  [save] {long_csv.name}  {len(rows)} rows "
      f"({len(ORDER)} contrasts x {len(LAYERS)} layers x {2*N_TOP})")

# compact L24 summaries
l24 = {}
for comp, table in per_contrast.items():
    a = comp.split("_minus_")[0]
    sel = [r for r in table if r["layer_ours"] == LOCKED]
    l24[comp] = {
        "positive_means_more_promoted_under": a,
        "promoted": [{"rank": r["rank"], "token": r["token_repr"], "id": r["token_id"],
                      "mean": r["mean_delta_logprob"], "median": r["median_delta_logprob"],
                      "sign_consistency": r["sign_consistency"]}
                     for r in sel if r["role"] == "promoted"],
        "suppressed": [{"rank": r["rank"], "token": r["token_repr"], "id": r["token_id"],
                        "mean": r["mean_delta_logprob"], "median": r["median_delta_logprob"],
                        "sign_consistency": r["sign_consistency"]}
                       for r in sel if r["role"] == "suppressed"],
    }
(OUT / "L24_summary.json").write_text(json.dumps(l24, indent=2, ensure_ascii=False),
                                      encoding="utf-8")
with (OUT / "L24_summary.csv").open("w", newline="", encoding="utf-8") as f:
    w = csv.DictWriter(f, fieldnames=["contrast", "role", "rank", "token_id", "token_repr",
                                      "mean_delta_logprob", "median_delta_logprob",
                                      "sign_consistency"])
    w.writeheader()
    for comp, table in per_contrast.items():
        for r in table:
            if r["layer_ours"] == LOCKED:
                w.writerow({k: r[k] for k in w.fieldnames})
print(f"  [save] L24_summary.json / L24_summary.csv")

# Excel workbook, one sheet per contrast
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    hdr = ["layer_ours", "layer_locked", "role", "rank", "token_id", "token_repr",
           "mean_delta_logprob", "median_delta_logprob", "sign_consistency"]
    for comp in ORDER:
        a, b = comp.split("_minus_")
        ws = wb.create_sheet(comp[:31])
        ws.append([f"{comp}   |   positive = MORE PROMOTED under '{a}'   |   "
                   f"negative = more promoted under '{b}'"])
        ws.append([f"Jacobian-Lens signed log-prob delta at final_input. n=72 tasks. "
                   f"L{LOCKED} is the locked layer; all others descriptive. {SUBWORD}"])
        ws.append([])
        ws.append(hdr)
        for c in ws[4]:
            c.font = Font(bold=True)
        for r in per_contrast[comp]:
            ws.append([r[k] for k in hdr])
        fill = PatternFill("solid", start_color="FFF2CC")
        for row in ws.iter_rows(min_row=5):
            if row[1].value is True:
                for c in row:
                    c.fill = fill
        widths = [11, 12, 11, 6, 9, 26, 20, 21, 17]
        for i, wdt in enumerate(widths, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = wdt
        ws.freeze_panes = "A5"
    xl = OUT / "jlens_layer_tables.xlsx"
    wb.save(xl)
    print(f"  [save] {xl.name}  ({len(ORDER)} sheets, L{LOCKED} rows highlighted)")
except ImportError:
    print("  openpyxl not installed -- workbook skipped, CSVs written")

(OUT / "PROVENANCE.json").write_text(json.dumps({
    "generated_from": "archived Jacobian delta tensors; no recomputation, no model, no GPU",
    "tensor_meta": {k: meta[k] for k in ("instrument", "is_logit_lens", "quantity",
                                         "sign_convention", "our_layers_covered",
                                         "vocab_size", "n_tasks", "locked_layer",
                                         "locked_window")},
    "ranking": {"promoted": f"top {N_TOP} by mean signed delta",
                "suppressed": f"bottom {N_TOP} by mean signed delta",
                "sign_consistency": "fraction of tasks matching the mean's sign",
                "manual_selection": False},
    "vocab_repair": {
        "problem": "the run-time dump used tok.vocab_size (128000) so ids 128000-128255 "
                   "were unmapped and decoded as <unk> in 2 of 600 earlier CSV rows",
        "fix": "added the 256 special/reserved token strings",
        "source": "locally cached NousResearch/Meta-Llama-3.1-8B-Instruct tokenizer, "
                  "loaded offline with clean_up_tokenization_spaces=False",
        "equivalence_proof": "all 128000 base-vocab ids decode byte-identically to the "
                             "archived dump; 0 mismatches",
        "note": "the initial comparison showed 117 mismatches, all a stripped leading space "
                "before punctuation -- transformers 4.49 applies clean_up_tokenization_spaces "
                "while the 5.15 run-time ignored it for BPE. Disabling the flag reproduces "
                "the archived strings exactly, so the ARCHIVED strings are the faithful ones.",
        "model_weights_downloaded": False, "gpu_used": False,
    },
}, indent=2), encoding="utf-8")
print(f"  [save] PROVENANCE.json")
